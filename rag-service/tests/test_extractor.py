import io
import unittest

from docx import Document
from openpyxl import Workbook

from services.extractor import extract_text, get_page_count


class ExtractorTests(unittest.TestCase):
    def test_extracts_word_paragraphs_and_tables(self):
        document = Document()
        document.add_heading("Inspection Summary", level=1)
        document.add_paragraph("Seal leakage requires corrective action.")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Equipment"
        table.cell(0, 1).text = "Risk"
        table.cell(1, 0).text = "Pump P-101"
        table.cell(1, 1).text = "High"
        stream = io.BytesIO()
        document.save(stream)

        text = extract_text(stream.getvalue(), "inspection.docx")

        self.assertIn("Inspection Summary", text)
        self.assertIn("Pump P-101 | High", text)
        self.assertEqual(1, get_page_count(stream.getvalue(), "inspection.docx"))

    def test_extracts_excel_cells_and_sheet_names(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Risk Register"
        sheet.append(["Activity", "Risk"])
        sheet.append(["Pump maintenance", "High"])
        stream = io.BytesIO()
        workbook.save(stream)

        text = extract_text(stream.getvalue(), "risk-register.xlsx")

        self.assertIn("Sheet: Risk Register", text)
        self.assertIn("Activity | Risk", text)
        self.assertIn("Pump maintenance | High", text)
        self.assertEqual(1, get_page_count(stream.getvalue(), "risk-register.xlsx"))

    def test_rejects_mismatched_unsupported_extension(self):
        with self.assertRaisesRegex(ValueError, "Unsupported file type"):
            extract_text(b"plain data", "legacy.xls")


if __name__ == "__main__":
    unittest.main()
