"""
Extracts plain text from uploaded files.
- Text PDFs  → pdfplumber (fast, accurate)
- Scanned PDFs / images → pdf2image + pytesseract (OCR)
- Word and Excel files → native document parsers
- Plain .txt files → direct read
"""
import io
import os
import shutil
import pytesseract
import pdfplumber
from pdf2image import convert_from_bytes
from PIL import Image
from docx import Document
from openpyxl import load_workbook


def _configure_tesseract() -> None:
    """Find Tesseract from configuration, PATH, or common Windows installs."""
    candidates = [
        os.getenv("TESSERACT_CMD"),
        shutil.which("tesseract"),
        r"C:\Program Files\Tesseract-OCR\tesseract.exe",
        r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
    ]

    for candidate in candidates:
        if not candidate:
            continue
        resolved = candidate if os.path.isfile(candidate) else shutil.which(candidate)
        if resolved:
            pytesseract.pytesseract.tesseract_cmd = resolved
            return


_configure_tesseract()


def _is_scanned_pdf(pdf_bytes: bytes) -> bool:
    """Heuristic: if pdfplumber extracts <50 chars across all pages, treat as scanned."""
    with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
        total_text = "".join(
            (page.extract_text() or "") for page in pdf.pages
        )
    return len(total_text.strip()) < 50


def extract_text(file_bytes: bytes, filename: str) -> str:
    ext = os.path.splitext(filename)[-1].lower()

    if ext == ".txt":
        return file_bytes.decode("utf-8", errors="replace")

    if ext == ".pdf":
        if _is_scanned_pdf(file_bytes):
            return _ocr_pdf(file_bytes)
        else:
            return _extract_text_pdf(file_bytes)

    if ext in {".png", ".jpg", ".jpeg", ".tiff", ".bmp"}:
        return _ocr_image(file_bytes)

    if ext == ".docx":
        return _extract_docx(file_bytes)

    if ext in {".xlsx", ".xlsm"}:
        return _extract_xlsx(file_bytes)

    raise ValueError(f"Unsupported file type: {ext}")


def get_page_count(file_bytes: bytes, filename: str) -> int:
    """Return a useful page/sheet count for document-library metadata."""
    ext = os.path.splitext(filename)[-1].lower()
    if ext == ".pdf":
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            return len(pdf.pages)
    if ext == ".docx":
        document = Document(io.BytesIO(file_bytes))
        return max(1, len(document.sections))
    if ext in {".xlsx", ".xlsm"}:
        workbook = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
        try:
            return max(1, len(workbook.sheetnames))
        finally:
            workbook.close()
    return 1


def _extract_text_pdf(pdf_bytes: bytes) -> str:
    pages = []
    with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
        for page in pdf.pages:
            pages.append(page.extract_text() or "")
    return "\n\n".join(pages)


def _ocr_pdf(pdf_bytes: bytes) -> str:
    images = convert_from_bytes(pdf_bytes, dpi=300)
    texts = [pytesseract.image_to_string(img) for img in images]
    return "\n\n".join(texts)


def _ocr_image(img_bytes: bytes) -> str:
    img = Image.open(io.BytesIO(img_bytes))
    return pytesseract.image_to_string(img)


def _extract_docx(docx_bytes: bytes) -> str:
    document = Document(io.BytesIO(docx_bytes))
    blocks = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells]
            if any(values):
                blocks.append(" | ".join(values))
    return "\n\n".join(blocks)


def _extract_xlsx(xlsx_bytes: bytes) -> str:
    workbook = load_workbook(io.BytesIO(xlsx_bytes), read_only=True, data_only=True)
    blocks = []
    try:
        for sheet in workbook.worksheets:
            blocks.append(f"Sheet: {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                values = ["" if value is None else str(value).strip() for value in row]
                while values and not values[-1]:
                    values.pop()
                if any(values):
                    blocks.append(" | ".join(values))
    finally:
        workbook.close()
    return "\n".join(blocks)
